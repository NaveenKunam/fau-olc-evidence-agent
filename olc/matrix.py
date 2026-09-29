"""Import the existing FAU OLC Evidence Matrix (xlsx) as PRIOR CLAIMS that need verification.

Matrix statements are summaries written by the team; they have no verbatim passage or page number. They are
therefore imported with ai_status = NEEDS VERIFICATION and strength = WEAK, and linked to passage-level evidence
from the cited source when the agent finds it ("corroborated_by").
"""
import json
import re

import openpyxl

from .analyze import fingerprint, next_code
from .db import now, audit

CAT_MAP = {"Institutional Support": "INS", "Technology Support": "TEC",
           "Course Development & Instructional Design": "CDID", "Course Structure": "CS",
           "Teaching & Learning": "TL", "Faculty Support": "FAC", "Learner Support": "LEA",
           "Evaluation & Assessment": "EVA"}


def canon(raw_id, category):
    code = CAT_MAP.get((category or "").strip())
    m = re.search(r"(\d+)", raw_id or "")
    if not code or not m:
        return None
    return f"{code}-{int(m.group(1)):02d}"


def _level_for(src):
    s = (src or "").lower()
    if "strategic-plan" in s or "strategic plan" in s:
        return "INTENT", "PLAN"
    if "polic" in s or "12-7" in s or "12-2" in s:
        return "PROCESS", "POLICY"
    return "PROCESS", "WEBPAGE"


def import_matrix(conn, path, actor="seed"):
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["Evidence Matrix"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h or "").strip() for h in rows[0]]
    col = {h: i for i, h in enumerate(hdr)}
    created, skipped = 0, 0
    doc = conn.execute("SELECT id FROM documents WHERE origin='matrix'").fetchone()
    if doc:
        doc_id = doc["id"]
    else:
        doc_id = conn.execute(
            """INSERT INTO documents(title, kind, doc_type, source_org, filename, authority, classification, origin, notes,
               added_by, added_at, processed_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
            ("FAU OLC Evidence Matrix (prior team analysis)", "matrix", "MATRIX", "Florida Atlantic Online",
             "FAU_OLC_Evidence_Matrix.xlsx", "FAU", "INTERNAL", "matrix",
             "Team-authored summaries. Not primary evidence; every row must be verified against its cited source.",
             actor, now(), now())).lastrowid
    for rn, r in enumerate(rows[1:], start=2):
        ind = canon(r[col["ID"]], r[col["Category"]])
        found = r[col["Evidence Found"]]
        if not ind:
            continue
        if not found:
            skipped += 1
            continue
        fp = fingerprint(ind, "matrix:" + str(found))
        if conn.execute("SELECT 1 FROM evidence WHERE fingerprint=?", (fp,)).fetchone():
            continue
        src = r[col["Source URL"]] or ""
        level, dtype = _level_for(src)
        prior = {"matrix_row": rn, "matrix_id": r[col["ID"]], "matrix_strength": r[col["Evidence Strength"]],
                 "matrix_missing": r[col["What Is Missing / Next Evidence"]], "matrix_owner": r[col["Likely Owner"]],
                 "matrix_indicator_wording": r[col["OLC Indicator"]], "source": src}
        conn.execute(
            """INSERT INTO evidence(evidence_code, indicator_id, document_id, title, evidence_type, source_org, source_ref,
               page, section, passage, summary, mapping_rationale, implementation_level, level_rationale, match_score, ai_status,
               strength, missing_note, next_artifact, likely_owner, origin, prior_claim_json, workflow_stage, classification,
               fingerprint, added_by, added_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (next_code(conn), ind, doc_id, f"Prior matrix claim (row {rn})", "Prior matrix claim", "Florida Atlantic Online",
             src, f"row {rn}", "Evidence Matrix", None, str(found),
             f"Imported from the FAU OLC Evidence Matrix row {rn} (matrix ID {r[col['ID']]}, rated '{r[col['Evidence Strength']]}' by the team). "
             "No verbatim passage or page was recorded, so this is a claim to verify, not evidence.",
             level, f"Cited source looks like a {dtype.lower()}; a claim based on it can reach {level} at most.", 0,
             "NEEDS VERIFICATION", "WEAK", r[col["What Is Missing / Next Evidence"]], None, r[col["Likely Owner"]], "matrix",
             json.dumps(prior, ensure_ascii=False), "NEEDS HUMAN REVIEW", "INTERNAL", fp, actor, now()))
        created += 1
    # evidence requests sheet
    reqs = 0
    if "Evidence Requests" in wb.sheetnames:
        for rn, r in enumerate(list(wb["Evidence Requests"].iter_rows(values_only=True))[1:], start=2):
            if not r or not r[0]:
                continue
            ind = canon(r[0], r[1])
            text = r[2] or ""
            if not ind or "No public evidence seeded yet" in text:
                continue
            if conn.execute("SELECT 1 FROM evidence_requests WHERE indicator_id=? AND request=?", (ind, text)).fetchone():
                continue
            conn.execute("""INSERT INTO evidence_requests(indicator_id, request, owner, status, assigned_to, notes, source, created_at)
                            VALUES (?,?,?,?,?,?,?,?)""", (ind, text, r[3], r[4] or "Open", r[5], r[6],
                                                          f"FAU matrix Evidence Requests row {rn}", now()))
            reqs += 1
    audit(conn, actor, "import_matrix", path, {"claims": created, "rows_without_evidence": skipped, "requests": reqs})
    conn.commit()
    return {"claims": created, "rows_without_evidence": skipped, "requests": reqs}
