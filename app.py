"""FAU OLC Evidence Agent - web application."""
import io
import json
import os
import secrets
from functools import wraps

from flask import (Flask, Response, abort, flash, g, redirect, render_template, request, send_file, session, url_for)
from markupsafe import Markup
from werkzeug.security import check_password_hash, generate_password_hash

from olc import commands, reports
from olc import ingest, md, rubric
from olc.analyze import process_document
from olc.assess import assess_all, assess_one, display_status, load_evidence, summary, STRENGTH_RANK
from olc.connectors import connectors
from olc.db import CLEARANCE_RANK, audit, close_db, connect, ensure_dirs, get_db, init_schema, now, visible_classes
from olc.indicators import LEVEL_HELP, LEVELS, RUBRIC, all_indicators, get_indicator, normalize_id, update_indicator_text
from olc.llm import get_llm
from olc.service import add_bytes, add_url, checkpoint, search_fau

BASE = os.path.dirname(os.path.abspath(__file__))
ROLE_RANK = {"viewer": 0, "contributor": 1, "reviewer": 2, "admin": 3}


def create_app():
    app = Flask(__name__, instance_path=os.environ.get("OLC_INSTANCE", os.path.join(BASE, "instance")))
    ensure_dirs(app.instance_path)
    secret_file = os.path.join(app.instance_path, "secret_key")
    if not os.path.exists(secret_file):
        with open(secret_file, "w") as f:
            f.write(secrets.token_hex(32))
    app.config.update(
        SECRET_KEY=os.environ.get("OLC_SECRET_KEY") or open(secret_file).read().strip(),
        DATABASE=os.path.join(app.instance_path, "olc.db"),
        UPLOADS=os.path.join(app.instance_path, "uploads"),
        SOURCES_CONFIG=os.path.join(BASE, "config", "approved_sources.json"),
        MAX_CONTENT_LENGTH=40 * 1024 * 1024,
        SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=bool(os.environ.get("OLC_SECURE_COOKIES")),
    )
    conn = connect(app.config["DATABASE"])
    init_schema(conn)
    conn.close()
    app.teardown_appcontext(close_db)
    register(app)
    return app


# ------------------------------------------------------------------ helpers

def cfg():
    with open(os.path.join(BASE, "config", "approved_sources.json"), encoding="utf-8") as f:
        return json.load(f)


def current_user():
    uid = session.get("uid")
    if not uid:
        return None
    if "user" not in g:
        r = get_db().execute("SELECT * FROM users WHERE id=? AND active=1", (uid,)).fetchone()
        g.user = dict(r) if r else None
    return g.user


def login_required(role="viewer"):
    def deco(fn):
        @wraps(fn)
        def wrapper(*a, **kw):
            u = current_user()
            if not u:
                return redirect(url_for("login", next=request.path))
            if ROLE_RANK[u["role"]] < ROLE_RANK[role]:
                abort(403)
            if request.method == "POST" and request.form.get("csrf") != session.get("csrf"):
                abort(400, "Form expired. Reload the page and try again.")
            return fn(*a, **kw)
        return wrapper
    return deco


def can(role):
    u = current_user()
    return bool(u and ROLE_RANK[u["role"]] >= ROLE_RANK[role])


def clearance():
    u = current_user()
    return u["clearance"] if u else "PUBLIC"


def actor():
    u = current_user()
    return u["display_name"] if u else "system"


def register(app):
    @app.context_processor
    def inject():
        if "csrf" not in session:
            session["csrf"] = secrets.token_hex(16)
        return dict(user=current_user(), can=can, csrf=session["csrf"], LEVELS=LEVELS, LEVEL_HELP=LEVEL_HELP,
                    RUBRIC=RUBRIC, llm_on=get_llm() is not None, display_status=display_status)

    @app.template_filter("md")
    def md_filter(s):
        return Markup(md.render(s))

    @app.template_filter("slug")
    def slug(s):
        return (s or "").lower().replace("/", "-").replace(" ", "-")

    # -------------------------------------------------------------- auth
    @app.route("/login", methods=["GET", "POST"])
    def login():
        if request.method == "POST":
            db = get_db()
            r = db.execute("SELECT * FROM users WHERE username=? AND active=1", (request.form.get("username", "").strip().lower(),)).fetchone()
            if r and check_password_hash(r["pw_hash"], request.form.get("password", "")):
                session.clear()
                session["uid"] = r["id"]
                session["csrf"] = secrets.token_hex(16)
                audit(db, r["display_name"], "login")
                db.commit()
                nxt = request.args.get("next") or "/"
                return redirect(nxt if nxt.startswith("/") else "/")
            flash("Invalid username or password.", "error")
        return render_template("login.html")

    @app.route("/logout")
    def logout():
        session.clear()
        return redirect(url_for("login"))

    # -------------------------------------------------------------- dashboard
    @app.route("/")
    @login_required()
    def dashboard():
        res = assess_all(get_db(), clearance())
        s = summary(res)
        f = {k: request.args.get(k, "") for k in ("category", "status", "strength", "owner", "review", "q")}
        rows = []
        for x in res:
            i = x["indicator"]
            if f["category"] and i["category_code"] != f["category"]:
                continue
            if f["status"] and x["status"] != f["status"]:
                continue
            if f["strength"] and x["strength"] != f["strength"]:
                continue
            if f["owner"] and not any(f["owner"].lower() in o.lower() for o in x["owners"]):
                continue
            if f["review"] and x["review_status"] != f["review"]:
                continue
            if f["q"] and f["q"].lower() not in (i["id"] + " " + (i["text"] or "")).lower():
                continue
            rows.append(x)
        cats = []
        for x in res:
            c = x["indicator"]["category"]
            if not cats or cats[-1]["name"] != c:
                cats.append({"name": c, "code": x["indicator"]["category_code"], "inds": []})
            cats[-1]["inds"].append(x)
        owners = sorted({o for x in res for o in x["owners"]})
        return render_template("dashboard.html", s=s, rows=rows, cats=cats, f=f, owners=owners)

    # -------------------------------------------------------------- indicator
    @app.route("/indicator/<ind_id>")
    @login_required()
    def indicator(ind_id):
        iid = normalize_id(ind_id)
        if iid != ind_id:
            return redirect(url_for("indicator", ind_id=iid))
        x, ev = assess_one(get_db(), clearance(), iid)
        if not x:
            abort(404)
        reqs = [dict(r) for r in get_db().execute("SELECT * FROM evidence_requests WHERE indicator_id=? ORDER BY id", (iid,))]
        hist = [dict(r) for r in get_db().execute(
            "SELECT * FROM audit_log WHERE target=? OR target LIKE ? ORDER BY id DESC LIMIT 25", (f"ind:{iid}", f"ev:%:{iid}"))]
        all_ids = [i["id"] for i in all_indicators(get_db())]
        pos = all_ids.index(iid)
        nav = (all_ids[pos - 1] if pos else None, all_ids[pos + 1] if pos + 1 < len(all_ids) else None)
        return render_template("indicator.html", x=x, ev=ev, reqs=reqs, hist=hist, nav=nav, all_ids=all_ids)

    @app.route("/indicator/<ind_id>/review", methods=["POST"])
    @login_required("reviewer")
    def indicator_review(ind_id):
        db = get_db()
        iid = normalize_id(ind_id)
        f = request.form

        def score(v):
            return int(v) if v not in (None, "", "none") else None
        vals = (f.get("review_status", "NOT REVIEWED"), score(f.get("human_prelim_score")), score(f.get("official_score")),
                actor(), f.get("notes", "").strip() or None, now())
        if vals[2] is not None and vals[0] != "APPROVED":
            flash("Official score recorded. Review status set to APPROVED because an official score was entered.", "info")
            vals = ("APPROVED",) + vals[1:]
        db.execute("""INSERT INTO indicator_reviews(indicator_id, review_status, human_prelim_score, official_score, reviewer, notes, updated_at)
                      VALUES (?,?,?,?,?,?,?) ON CONFLICT(indicator_id) DO UPDATE SET review_status=excluded.review_status,
                      human_prelim_score=excluded.human_prelim_score, official_score=excluded.official_score, reviewer=excluded.reviewer,
                      notes=excluded.notes, updated_at=excluded.updated_at""", (iid,) + vals)
        audit(db, actor(), "indicator_review", f"ind:{iid}",
              {"status": vals[0], "human_prelim": vals[1], "official": vals[2], "notes": vals[4]})
        db.commit()
        flash("Reviewer decision saved.", "ok")
        return redirect(url_for("indicator", ind_id=iid) + "#review")

    @app.route("/indicator/<ind_id>/request", methods=["POST"])
    @login_required("contributor")
    def indicator_request(ind_id):
        db = get_db()
        iid = normalize_id(ind_id)
        if request.form.get("request_id"):
            db.execute("UPDATE evidence_requests SET status=?, assigned_to=?, notes=? WHERE id=?",
                       (request.form.get("status"), request.form.get("assigned_to"), request.form.get("notes"), request.form["request_id"]))
        else:
            db.execute("INSERT INTO evidence_requests(indicator_id, request, owner, status, source, created_at) VALUES (?,?,?,?,?,?)",
                       (iid, request.form["request"], request.form.get("owner"), "Open", f"Added by {actor()}", now()))
        audit(db, actor(), "evidence_request", f"ind:{iid}", dict(request.form))
        db.commit()
        return redirect(url_for("indicator", ind_id=iid) + "#requests")

    @app.route("/indicator/<ind_id>/evidence/new", methods=["POST"])
    @login_required("contributor")
    def evidence_new(ind_id):
        """Manual evidence entry: a reviewer records a passage they located themselves (must cite a source)."""
        db = get_db()
        iid = normalize_id(ind_id)
        f = request.form
        if not f.get("passage", "").strip() or not f.get("source_ref", "").strip():
            flash("Manual evidence needs a verbatim passage and a source (URL or filename). No evidence = no claim.", "error")
            return redirect(url_for("indicator", ind_id=iid))
        from olc.analyze import next_code, fingerprint, gap_note
        ind = get_indicator(db, iid)
        level = f.get("level") or "PROCESS"
        missing, nxt = gap_note(level, ind)
        code = next_code(db)
        db.execute("""INSERT INTO evidence(evidence_code, indicator_id, title, evidence_type, source_org, source_ref, doc_date, page, section,
                      passage, summary, mapping_rationale, implementation_level, level_rationale, match_score, ai_status, strength,
                      missing_note, next_artifact, likely_owner, origin, workflow_stage, classification, human_edited, fingerprint,
                      added_by, added_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                   (code, iid, f.get("title") or f["source_ref"], "Human-entered", f.get("source_org") or "Florida Atlantic University",
                    f["source_ref"], f.get("doc_date"), f.get("page"), f.get("section"), f["passage"].strip(),
                    f.get("summary") or f["passage"][:240], f.get("rationale") or "Mapped by a human reviewer.", level,
                    "Level assigned by the human reviewer who entered this evidence.", None, "NEEDS VERIFICATION",
                    f.get("strength") or "MODERATE", missing, nxt, "; ".join(ind["config"].get("owners") or []), "human",
                    "NEEDS HUMAN REVIEW", f.get("classification") or "INTERNAL", 1, fingerprint(iid, f["passage"]), actor(), now()))
        audit(db, actor(), "evidence_manual_add", f"ev:{code}:{iid}", {"source": f["source_ref"]})
        db.commit()
        flash(f"{code} added. It needs a second reviewer's approval before it counts as verified.", "ok")
        return redirect(url_for("indicator", ind_id=iid) + "#evidence")

    # -------------------------------------------------------------- evidence
    @app.route("/evidence")
    @login_required()
    def evidence():
        db = get_db()
        ev = load_evidence(db, clearance())
        f = {k: request.args.get(k, "") for k in ("indicator", "status", "strength", "review", "origin", "doc", "code", "q", "level")}
        out = []
        for e in ev:
            if f["indicator"] and e["indicator_id"] != normalize_id(f["indicator"]):
                continue
            if f["status"] and display_status(e) != f["status"]:
                continue
            if f["strength"] and e["strength"] != f["strength"]:
                continue
            if f["review"] and e["review_status"] != f["review"]:
                continue
            if f["origin"] and not e["origin"].startswith(f["origin"]):
                continue
            if f["doc"] and str(e["document_id"]) != f["doc"]:
                continue
            if f["code"] and e["evidence_code"] != f["code"]:
                continue
            if f["level"] and e["implementation_level"] != f["level"]:
                continue
            if f["q"] and f["q"].lower() not in ((e["passage"] or "") + (e["summary"] or "") + e["title"]).lower():
                continue
            out.append(e)
        out.sort(key=lambda e: e["id"])
        return render_template("evidence.html", ev=out, f=f, total=len(ev))

    @app.route("/evidence/<int:eid>", methods=["POST"])
    @login_required("contributor")
    def evidence_update(eid):
        db = get_db()
        e = db.execute("SELECT * FROM evidence WHERE id=?", (eid,)).fetchone()
        if not e or e["classification"] not in visible_classes(clearance()):
            abort(404)
        act = request.form.get("act")
        back = request.form.get("back") or url_for("indicator", ind_id=e["indicator_id"])
        reviewer_acts = {"approve", "reject", "include", "exclude", "reset"}
        if act in reviewer_acts and not can("reviewer"):
            abort(403)
        notes = request.form.get("reviewer_notes")
        if act == "approve":
            db.execute("""UPDATE evidence SET review_status='APPROVED', workflow_stage='APPROVED', reviewer=?, reviewed_at=?,
                          reviewer_notes=COALESCE(?, reviewer_notes), human_edited=1 WHERE id=?""", (actor(), now(), notes or None, eid))
        elif act == "reject":
            db.execute("""UPDATE evidence SET review_status='REJECTED', workflow_stage='REJECTED', include_in_submission=0, reviewer=?,
                          reviewed_at=?, reviewer_notes=COALESCE(?, reviewer_notes), human_edited=1 WHERE id=?""", (actor(), now(), notes or None, eid))
        elif act == "include":
            if e["review_status"] != "APPROVED":
                flash("Only APPROVED evidence can be included in the submission.", "error")
                return redirect(back)
            db.execute("UPDATE evidence SET include_in_submission=1, workflow_stage='INCLUDED IN SUBMISSION', human_edited=1 WHERE id=?", (eid,))
        elif act == "exclude":
            db.execute("UPDATE evidence SET include_in_submission=0, workflow_stage='APPROVED', human_edited=1 WHERE id=?", (eid,))
        elif act == "reset":
            db.execute("""UPDATE evidence SET review_status='NOT REVIEWED', workflow_stage='NEEDS HUMAN REVIEW', include_in_submission=0,
                          reviewer=?, reviewed_at=?, human_edited=1 WHERE id=?""", (actor(), now(), eid))
        elif act == "edit":
            f = request.form
            new_ind = normalize_id(f.get("indicator_id") or e["indicator_id"])
            if not get_indicator(db, new_ind):
                flash(f"Unknown indicator {new_ind}", "error")
                return redirect(back)
            cls = f.get("classification") or e["classification"]
            if CLEARANCE_RANK[cls] > CLEARANCE_RANK[clearance()]:
                flash("You cannot classify evidence above your own clearance.", "error")
                return redirect(back)
            db.execute("""UPDATE evidence SET indicator_id=?, implementation_level=?, strength=?, classification=?, reviewer_notes=?,
                          mapping_rationale=?, human_edited=1 WHERE id=?""",
                       (new_ind, f.get("implementation_level") or e["implementation_level"], f.get("strength") or e["strength"], cls,
                        f.get("reviewer_notes", e["reviewer_notes"]),
                        (f.get("mapping_rationale") or e["mapping_rationale"]) if f.get("mapping_rationale") is not None else e["mapping_rationale"], eid))
            back = request.form.get("back") or url_for("indicator", ind_id=new_ind)
        elif act == "also_map":
            new_ind = normalize_id(request.form.get("indicator_id", ""))
            if not get_indicator(db, new_ind):
                flash(f"Unknown indicator {new_ind}", "error")
                return redirect(back)
            from olc.analyze import next_code
            cols = [c for c in e.keys() if c not in ("id", "evidence_code", "indicator_id", "fingerprint")]
            row = dict(e)
            row.update(review_status="NOT REVIEWED", workflow_stage="NEEDS HUMAN REVIEW", include_in_submission=0, human_edited=1,
                       origin="human", added_by=actor(), added_at=now(), reviewer=None, reviewed_at=None,
                       mapping_rationale=f"Additional mapping by {actor()} from {e['evidence_code']}. " + (request.form.get("reason") or ""))
            code = next_code(db)
            db.execute(f"INSERT INTO evidence(evidence_code, indicator_id, {', '.join(cols)}) VALUES (?,?,{','.join('?' * len(cols))})",
                       [code, new_ind] + [row[c] for c in cols])
            flash(f"{code} created mapping this passage to {new_ind}.", "ok")
        else:
            abort(400)
        audit(db, actor(), f"evidence_{act}", f"ev:{e['evidence_code']}:{e['indicator_id']}", {"notes": notes} if notes else None)
        db.commit()
        return redirect(back)

    # -------------------------------------------------------------- sources
    @app.route("/sources")
    @login_required()
    def sources():
        db = get_db()
        cls = visible_classes(clearance())
        docs = [dict(r) for r in db.execute(
            f"""SELECT d.*, (SELECT COUNT(*) FROM evidence e WHERE e.document_id=d.id) AS n_ev,
                       (SELECT COUNT(*) FROM passages p WHERE p.document_id=d.id) AS n_pass
                FROM documents d WHERE d.classification IN ({','.join('?' * len(cls))}) ORDER BY d.id DESC""", cls)]
        return render_template("sources.html", docs=docs, conns=connectors(app.config["SOURCES_CONFIG"]))

    @app.route("/sources/add", methods=["GET", "POST"])
    @login_required("contributor")
    def add_source():
        db = get_db()
        result = None
        if request.method == "POST":
            mode = request.form.get("mode")
            cls = request.form.get("classification") or "INTERNAL"
            if CLEARANCE_RANK[cls] > CLEARANCE_RANK[clearance()]:
                flash("You cannot add a source above your own clearance.", "error")
                return redirect(url_for("add_source"))
            opts = dict(actor=actor(), classification=cls, authority=request.form.get("authority") or None,
                        doc_type=request.form.get("doc_type") or None, doc_date=request.form.get("doc_date") or None,
                        title=request.form.get("title") or None, llm=get_llm())
            try:
                if mode == "url":
                    result = add_url(db, request.form["url"].strip(), cfg()["allowed_web_domains"], **opts)
                elif mode == "text":
                    data = request.form.get("text", "").encode("utf-8")
                    result = add_bytes(db, data, filename=(request.form.get("title") or "pasted-text") + ".txt",
                                       upload_dir=app.config["UPLOADS"], **opts)
                else:
                    files = request.files.getlist("files")
                    results = []
                    for fl in files:
                        if not fl or not fl.filename:
                            continue
                        results.append(add_bytes(db, fl.read(), filename=fl.filename, ctype=fl.mimetype or "",
                                                 upload_dir=app.config["UPLOADS"], **opts))
                    if not results:
                        raise ValueError("Choose at least one file.")
                    result = results[0] if len(results) == 1 else {"multi": results}
            except (ValueError, PermissionError) as ex:
                flash(str(ex), "error")
                return redirect(url_for("add_source"))
            except Exception as ex:  # network etc.
                flash(f"Could not process source: {ex}", "error")
                return redirect(url_for("add_source"))
            items = result["multi"] if "multi" in result else [result]
            for r in items:
                if r["status"] == "duplicate":
                    flash(f"'{r['title']}' is already in the repository (identical file). Nothing changed.", "info")
                else:
                    flash(f"Processed '{r['title']}': {len(r['created'])} evidence item(s) mapped"
                          f"{' (new version: previous version marked superseded)' if r['status'] == 'new_version' else ''}.", "ok")
            if len(items) == 1:
                return redirect(url_for("source", doc_id=items[0]["doc_id"]))
            return redirect(url_for("sources"))
        return render_template("add_source.html", allowed=cfg()["allowed_web_domains"])

    @app.route("/source/<int:doc_id>")
    @login_required()
    def source(doc_id):
        db = get_db()
        d = db.execute("SELECT * FROM documents WHERE id=?", (doc_id,)).fetchone()
        if not d or d["classification"] not in visible_classes(clearance()):
            abort(404)
        ev = [e for e in load_evidence(db, clearance()) if e["document_id"] == doc_id]
        passages = [dict(r) for r in db.execute("SELECT * FROM passages WHERE document_id=? ORDER BY seq", (doc_id,))]
        mapped = {}
        for e in ev:
            if e["passage_id"]:
                mapped.setdefault(e["passage_id"], []).append(e)
        return render_template("source.html", d=dict(d), ev=ev, passages=passages, mapped=mapped,
                               show_all=request.args.get("all"))

    @app.route("/source/<int:doc_id>/file")
    @login_required()
    def source_file(doc_id):
        d = get_db().execute("SELECT * FROM documents WHERE id=?", (doc_id,)).fetchone()
        if not d or d["classification"] not in visible_classes(clearance()) or not d["stored_path"]:
            abort(404)
        path = os.path.realpath(d["stored_path"])
        if not path.startswith(os.path.realpath(app.config["UPLOADS"])) and not path.startswith(os.path.realpath(os.path.join(BASE, "seed"))):
            abort(404)
        return send_file(path, download_name=d["filename"] or os.path.basename(path))

    @app.route("/source/<int:doc_id>/update", methods=["POST"])
    @login_required("contributor")
    def source_update(doc_id):
        db = get_db()
        d = db.execute("SELECT * FROM documents WHERE id=?", (doc_id,)).fetchone()
        if not d or d["classification"] not in visible_classes(clearance()):
            abort(404)
        act = request.form.get("act")
        if act == "reprocess":
            r = process_document(db, doc_id, actor(), llm=get_llm())
            flash(f"Reprocessed: {len(r['created'])} new/updated AI mappings; {r['removed']} unreviewed AI mappings replaced. "
                  "Human-reviewed evidence was not changed.", "ok")
        elif act == "meta":
            f = request.form
            cls = f.get("classification") or d["classification"]
            if CLEARANCE_RANK[cls] > CLEARANCE_RANK[clearance()]:
                abort(403)
            db.execute("""UPDATE documents SET title=?, doc_type=?, doc_date=?, is_draft=?, version_note=?, authority=?, classification=?,
                          source_org=?, notes=? WHERE id=?""",
                       (f.get("title") or d["title"], f.get("doc_type") or d["doc_type"], f.get("doc_date") or None,
                        1 if f.get("is_draft") else 0, f.get("version_note") or None, f.get("authority") or d["authority"], cls,
                        f.get("source_org") or None, f.get("notes") or None, doc_id))
            db.execute("UPDATE evidence SET classification=? WHERE document_id=? AND human_edited=0", (cls, doc_id))
            audit(db, actor(), "source_meta", f"doc:{doc_id}", dict(f))
            db.commit()
            if f.get("reprocess"):
                process_document(db, doc_id, actor(), llm=get_llm())
            flash("Source metadata saved." + (" Unreviewed mappings recalculated." if f.get("reprocess") else ""), "ok")
        db.commit()
        return redirect(url_for("source", doc_id=doc_id))

    # -------------------------------------------------------------- commands & reports
    @app.route("/commands", methods=["GET", "POST"])
    @login_required()
    def console():
        db = get_db()
        out, link, results, cmd = None, None, None, request.values.get("cmd", "")
        if cmd:
            r = commands.run(cmd, dict(conn=db, clearance=clearance(), llm=get_llm()))
            act = r.get("action")
            if act and not can("contributor"):
                abort(403)
            if act and request.method != "POST":
                out = f"Press **Run** to execute `{cmd}`."
            elif act == "process_url":
                try:
                    res = add_url(db, r["url"], cfg()["allowed_web_domains"], actor=actor(), llm=get_llm())
                    if res["status"] == "duplicate":
                        out = f"Already in the repository: **{res['title']}** (source #{res['doc_id']})."
                    else:
                        out = (f"# PROCESS SOURCE\nProcessed **{res['title']}** (source #{res['doc_id']}, {res['passages']} passages).\n\n"
                               f"{len(res['created'])} evidence item(s) mapped:\n" +
                               "\n".join(f"- {c} -> [{i}](/indicator/{i})" for c, i in res["created"]))
                    link = f"/source/{res['doc_id']}"
                except Exception as ex:
                    out = f"**Could not process:** {ex}"
            elif act == "reprocess":
                d = db.execute("SELECT classification FROM documents WHERE id=?", (r["doc_id"],)).fetchone()
                if not d or d["classification"] not in visible_classes(clearance()):
                    out = "Source not found."
                else:
                    res = process_document(db, r["doc_id"], actor(), llm=get_llm())
                    out = f"Reprocessed source #{r['doc_id']}: {len(res['created'])} mappings created, {res['removed']} unreviewed AI mappings replaced."
            elif act == "checkpoint":
                checkpoint(db, r["label"], actor(), clearance())
                out = f"Checkpoint **{r['label']}** saved. WHAT CHANGED will compare against it."
            elif act == "search":
                if not r["indicator"]:
                    out = "Usage: `SEARCH FAU <indicator>` e.g. `SEARCH FAU TEC-01`"
                else:
                    ind = get_indicator(db, r["indicator"])
                    if not ind or ind["pending"]:
                        out = "Unknown indicator, or its wording is pending (search terms come from the wording)."
                    else:
                        c = cfg()
                        results, n = search_fau(db, ind, c["allowed_web_domains"], c["fau_seed_pages"], max_pages=25)
                        out = (f"# SEARCH FAU {ind['id']}\nCrawled {n} approved FAU pages (bounded, read-only). "
                               f"{len(results)} candidate page(s) scored as relevant. Relevance is not substantiation: "
                               "process a candidate to extract and grade passages.")
            else:
                out = r.get("markdown")
                link = r.get("link")
            db.commit()
        return render_template("commands.html", out=out, cmd=cmd, link=link, results=results, ind=locals().get("ind"))

    @app.route("/reports/<name>")
    @login_required()
    def report(name):
        db = get_db()
        cl = clearance()
        fn = {"initial": lambda: reports.initial_review(db, cl, docs_note=missing_inputs(db)),
              "gaps": lambda: reports.gap_report(db, cl),
              "meeting": lambda: reports.meeting_brief(db, cl),
              "changed": lambda: reports.what_changed(db, cl),
              "verify": lambda: reports.verify_report(db, cl)}.get(name)
        if not fn:
            abort(404)
        text = fn()
        if request.args.get("download"):
            return Response(text, mimetype="text/markdown",
                            headers={"Content-Disposition": f"attachment; filename=fau-olc-{name}.md"})
        return render_template("report.html", text=text, name=name)

    @app.route("/owners")
    @login_required()
    def owners():
        res = assess_all(get_db(), clearance())
        ow = sorted(reports.all_owners(res).items(), key=lambda kv: -len(kv[1]))
        key = request.args.get("owner")
        packet = reports.owner_packet(get_db(), clearance(), key) if key else None
        if key and request.args.get("download"):
            return Response(packet, mimetype="text/markdown",
                            headers={"Content-Disposition": f"attachment; filename=olc-request-{key[:30].replace(' ', '_')}.md"})
        return render_template("owners.html", ow=ow, packet=packet, key=key)

    @app.route("/submission/<ind_id>")
    @login_required()
    def submission(ind_id):
        text = reports.submission_draft(get_db(), clearance(), normalize_id(ind_id), llm=get_llm())
        if request.args.get("download"):
            return Response(text, mimetype="text/markdown",
                            headers={"Content-Disposition": f"attachment; filename=olc-submission-{normalize_id(ind_id)}.md"})
        return render_template("report.html", text=text, name=f"submission-{ind_id}")

    @app.route("/search", methods=["POST"])
    @login_required("contributor")
    def search():
        db = get_db()
        ind = get_indicator(db, normalize_id(request.form["indicator"]))
        c = cfg()
        results, n = search_fau(db, ind, c["allowed_web_domains"], c["fau_seed_pages"], max_pages=25)
        return render_template("search.html", ind=ind, results=results, n=n)

    @app.route("/export.xlsx")
    @login_required()
    def export():
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        db = get_db()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Indicators"
        hdr = ["Indicator", "Category", "OLC Requirement", "Wording source", "Status", "Strength", "Has level", "Needs level",
               "AI Preliminary (advisory)", "AI Rationale", "Human preliminary", "Official score", "Review status", "Reviewer notes",
               "What is missing", "Recommended artifacts", "Likely owner", "Skeptical-review flags"]
        ws.append(hdr)
        for x in assess_all(db, clearance()):
            i = x["indicator"]
            ws.append([i["id"], i["category"], i["text"] or "PENDING", i["text_source"], x["status"], x["strength"], x["achieved"],
                       x["required"], x["prelim"], x["prelim_rationale"], x["human_prelim"], x["official"], x["review_status"],
                       x["review"].get("notes"), x["headline"], "\n".join(f"[{l}] {t}" for l, t in x["needed"]), ", ".join(x["owners"]),
                       "\n".join(f"[{f['severity']}] {f['message']}" for f in x["flags"])])
        ws2 = wb.create_sheet("Evidence")
        cols = ["Evidence ID", "OLC Indicator ID", "OLC Category", "OLC Requirement", "Evidence Title", "Evidence Type", "Source Organization",
                "Source URL or Filename", "Document Date", "Page Number", "Section", "Exact Supporting Passage", "Evidence Summary",
                "Why This Evidence Maps", "Evidence Status", "Evidence Strength", "Implementation Level", "Missing Evidence",
                "Recommended Next Artifact", "Likely FAU Owner", "Human Review Status", "Workflow Stage", "In Submission",
                "Classification", "Reviewer", "Reviewer Notes", "Added By", "Date Added", "Date Reviewed"]
        ws2.append(cols)
        inds = {i["id"]: i for i in all_indicators(db)}
        for e in load_evidence(db, clearance()):
            i = inds[e["indicator_id"]]
            ws2.append([e["evidence_code"], e["indicator_id"], i["category"], i["text"], e["title"], e["evidence_type"], e["source_org"],
                        e["doc_url"] or e["source_ref"], e["doc_date"], e["page"], e["section"], e["passage"], e["summary"],
                        e["mapping_rationale"], display_status(e), e["strength"], e["implementation_level"], e["missing_note"],
                        e["next_artifact"], e["likely_owner"], e["review_status"], e["workflow_stage"], "Yes" if e["include_in_submission"] else "",
                        e["classification"], e["reviewer"], e["reviewer_notes"], e["added_by"], e["added_at"], e["reviewed_at"]])
        for sh in (ws, ws2):
            for c in sh[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="003366")
            for colc in sh.columns:
                sh.column_dimensions[colc[0].column_letter].width = 28
            for row in sh.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            sh.freeze_panes = "B2"
        ws3 = wb.create_sheet("Read me")
        ws3.append(["AI preliminary scores are advisory only. Official FAU scores are entered by human reviewers."])
        ws3.append([f"Exported {now()} by {actor()} (clearance {clearance()}). Contains only evidence visible at this clearance."])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        audit(db, actor(), "export_xlsx")
        db.commit()
        return send_file(buf, download_name="FAU_OLC_Evidence_Repository.xlsx", as_attachment=True,
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    # -------------------------------------------------------------- admin
    @app.route("/admin")
    @login_required("admin")
    def admin():
        db = get_db()
        users = [dict(r) for r in db.execute("SELECT * FROM users ORDER BY id")]
        runs = [dict(r) for r in db.execute("SELECT id, label, created_at, created_by FROM runs ORDER BY id DESC LIMIT 10")]
        log = [dict(r) for r in db.execute("SELECT * FROM audit_log ORDER BY id DESC LIMIT 60")]
        inds = all_indicators(db)
        return render_template("admin.html", users=users, runs=runs, log=log, inds=inds, n_qss=sum(1 for i in inds if i["text_source"].startswith("OLC QSS")),
                               conns=connectors(app.config["SOURCES_CONFIG"]), c=cfg())

    @app.route("/admin/users", methods=["POST"])
    @login_required("admin")
    def admin_users():
        db = get_db()
        f = request.form
        if f.get("act") == "create":
            uname = f["username"].strip().lower()
            if not uname or len(f.get("password", "")) < 10:
                flash("Username required; password must be at least 10 characters.", "error")
                return redirect(url_for("admin"))
            try:
                db.execute("INSERT INTO users(username, display_name, role, clearance, pw_hash, created_at) VALUES (?,?,?,?,?,?)",
                           (uname, f.get("display_name") or uname, f.get("role", "reviewer"), f.get("clearance", "INTERNAL"),
                            generate_password_hash(f["password"]), now()))
            except Exception:
                flash("That username already exists.", "error")
                return redirect(url_for("admin"))
            audit(db, actor(), "user_create", uname, {"role": f.get("role"), "clearance": f.get("clearance")})
        elif f.get("act") == "update":
            uid = int(f["uid"])
            if uid == current_user()["id"] and (f.get("role") != "admin" or not f.get("active")):
                flash("You cannot demote or deactivate yourself.", "error")
                return redirect(url_for("admin"))
            db.execute("UPDATE users SET role=?, clearance=?, active=? WHERE id=?",
                       (f.get("role"), f.get("clearance"), 1 if f.get("active") else 0, uid))
            if f.get("password"):
                if len(f["password"]) < 10:
                    flash("Password must be at least 10 characters.", "error")
                    return redirect(url_for("admin"))
                db.execute("UPDATE users SET pw_hash=? WHERE id=?", (generate_password_hash(f["password"]), uid))
            audit(db, actor(), "user_update", str(uid), {"role": f.get("role"), "clearance": f.get("clearance")})
        db.commit()
        flash("User settings saved.", "ok")
        return redirect(url_for("admin"))

    @app.route("/account", methods=["GET", "POST"])
    @login_required()
    def account():
        if request.method == "POST":
            db = get_db()
            u = current_user()
            if not check_password_hash(u["pw_hash"], request.form.get("current", "")):
                flash("Current password is incorrect.", "error")
            elif len(request.form.get("new", "")) < 10:
                flash("New password must be at least 10 characters.", "error")
            else:
                db.execute("UPDATE users SET pw_hash=? WHERE id=?", (generate_password_hash(request.form["new"]), u["id"]))
                audit(db, u["display_name"], "password_change")
                db.commit()
                flash("Password changed.", "ok")
        return render_template("account.html")

    @app.route("/admin/indicator/<ind_id>", methods=["POST"])
    @login_required("admin")
    def admin_indicator(ind_id):
        db = get_db()
        iid = normalize_id(ind_id)
        f = request.form
        update_indicator_text(db, iid, f.get("text"), f.get("text_source") or f"Entered by {actor()} {now()}", actor(),
                              alt_text=f.get("alt_text"), handbook=f.get("handbook_text"),
                              handbook_source=f.get("handbook_source") or f"Entered by {actor()}")
        audit(db, actor(), "indicator_edit", f"ind:{iid}", {"text": f.get("text")})
        db.commit()
        flash(f"{iid} wording saved. Reprocess sources to map evidence against the new wording.", "ok")
        return redirect(url_for("indicator", ind_id=iid))

    @app.route("/admin/rubric", methods=["POST"])
    @login_required("admin")
    def admin_rubric():
        db = get_db()
        kind = request.form.get("kind")
        if request.form.get("apply"):
            staged = session.pop("rubric_stage", None)
            if not staged:
                flash("Nothing staged. Upload the PDF again.", "error")
                return redirect(url_for("admin"))
            with open(staged["path"], "rb") as fh:
                data = fh.read()
            if staged["kind"] == "qss":
                parsed = rubric.parse_qss(data)
                ids = set(request.form.getlist("ids"))
                n = rubric.apply_qss(db, parsed, staged["name"], actor(), ids=ids)
            else:
                parsed = rubric.parse_handbook(data, all_indicators(db))
                n = rubric.apply_handbook(db, parsed, staged["name"], actor())
            flash(f"Applied {n} indicator update(s) from {staged['name']}. Reprocess sources to re-map against new wording.", "ok")
            return redirect(url_for("admin"))
        fl = request.files.get("pdf")
        if not fl or not fl.filename.lower().endswith(".pdf"):
            flash("Upload the OLC PDF.", "error")
            return redirect(url_for("admin"))
        data = fl.read()
        path = os.path.join(app.config["UPLOADS"], f"rubric_{ingest.sha256(data)[:10]}.pdf")
        with open(path, "wb") as fh:
            fh.write(data)
        parsed = rubric.parse_qss(data) if kind == "qss" else rubric.parse_handbook(data, all_indicators(db))
        session["rubric_stage"] = {"path": path, "name": fl.filename, "kind": kind}
        rows = rubric.preview(db, parsed)
        return render_template("rubric_preview.html", rows=rows, kind=kind, name=fl.filename, n_found=len(parsed))

    @app.route("/admin/checkpoint", methods=["POST"])
    @login_required("reviewer")
    def admin_checkpoint():
        checkpoint(get_db(), request.form.get("label") or "Manual checkpoint", actor(), clearance())
        flash("Checkpoint saved.", "ok")
        return redirect(request.form.get("back") or url_for("dashboard"))

    @app.route("/connect/<provider>")
    @login_required("contributor")
    def connect_provider(provider):
        c = {x.name.split()[0].lower(): x for x in connectors(app.config["SOURCES_CONFIG"])}
        conn = c.get("google" if provider == "google" else "microsoft")
        flash(f"{conn.name if conn else provider}: {conn.status() if conn else 'unknown provider'}. "
              "See README > Enabling Google Drive / SharePoint.", "info")
        return redirect(url_for("sources"))

    @app.errorhandler(403)
    def forbidden(_e):
        return render_template("error.html", code=403, msg="Your role does not allow this action."), 403

    @app.errorhandler(404)
    def notfound(_e):
        return render_template("error.html", code=404, msg="Not found (or not visible at your clearance)."), 404


def missing_inputs(db):
    notes = []
    names = " ".join((r["filename"] or "") + " " + (r["title"] or "") for r in db.execute("SELECT filename, title FROM documents")).lower()
    if not any(i["text_source"].startswith("OLC QSS") for i in all_indicators(db)):
        notes.append("**QSS - Administration of Online Programs.pdf** was not supplied to this build. Indicator wording comes from the FAU "
                     "matrix; 13 slots have no wording. Admin > OLC Rubric > import the QSS PDF.")
    if not any(i.get("handbook_text") for i in all_indicators(db)):
        notes.append("**Administration of Online Programs Handbook.pdf** was not supplied. Handbook interpretation fields are empty until it is imported.")
    if "12-7" not in names and "12.7" not in names:
        notes.append("**FAU Policy 12.7 (System and Data Classifications)**: cited in the matrix as an upload but not supplied. Upload it to verify TEC-03/TEC-06 claims.")
    if "12-2" not in names and "12.2" not in names:
        notes.append("**FAU Policy 12.2 (Acceptable Use of Technology Resources)**: cited in the matrix but not supplied. Upload it to verify TEC-03 claims.")
    return notes


app = create_app()

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5210))
    app.run(host=os.environ.get("HOST", "127.0.0.1"), port=port, debug=bool(os.environ.get("OLC_DEBUG")))
